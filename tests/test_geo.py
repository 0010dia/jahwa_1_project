from io import BytesIO

import app as app_module
import pytest
import shapely
from admin_districts import get_admin_district_repository
from openpyxl import load_workbook
from app import (
    _extract_tmap_route,
    detour_waypoint_candidates,
    detour_waypoints,
    hazards_for_path,
    hazards_for_segment,
    haversine_m,
)
from flood_data import (
    FloodRepository,
    get_flood_repository,
    live_flood_geometry,
    load_flood_display,
)
from scripts.fetch_recent_flood_traces import record_matches
from scripts.fetch_gumi_shelters import is_gumi_temporary_housing


@pytest.fixture
def route_shelter(monkeypatch):
    shelter = {
        "id": "test-shelter",
        "name": "테스트 임시주거시설",
        "type": "이재민 임시주거시설",
        "address": "경상북도 구미시 테스트로 1",
        "lat": 36.11953,
        "lng": 128.34474,
        "capacity": 100,
        "verified": True,
    }
    monkeypatch.setattr(app_module, "load_shelters", lambda: [shelter])
    return shelter


@pytest.fixture
def report_db(tmp_path, monkeypatch):
    database_path = tmp_path / "reports.sqlite3"
    monkeypatch.setattr(app_module, "REPORTS_DB_PATH", database_path)
    return database_path


@pytest.fixture
def fake_flood_repository(monkeypatch):
    geometry = shapely.Polygon(
        [
            (128.336, 36.122),
            (128.340, 36.122),
            (128.340, 36.127),
            (128.336, 36.127),
        ]
    )
    zone = {
        "id": "flood-demo-001",
        "sourceId": "test",
        "name": "테스트 침수구역",
        "riskLevel": "테스트",
        "note": "테스트용 도형",
        "detour_margin_deg": 0.002,
        "bounds": list(geometry.bounds),
        "_geometry": geometry,
    }

    class FakeRepository:
        @staticmethod
        def hazards_for_paths(paths):
            for path in paths:
                if len(path) < 2:
                    continue
                line = shapely.LineString(
                    [(point["lng"], point["lat"]) for point in path]
                )
                if line.intersects(geometry):
                    return [zone]
            return []

        @staticmethod
        def hazards_for_point(point):
            return [zone] if geometry.intersects(
                shapely.Point(point["lng"], point["lat"])
            ) else []

        @staticmethod
        def contains_point(point):
            return bool(FakeRepository.hazards_for_point(point))

        @staticmethod
        def segment_intersects_zone(start, end, candidate_zone):
            return shapely.LineString(
                [(start["lng"], start["lat"]), (end["lng"], end["lat"])]
            ).intersects(candidate_zone["_geometry"])

    repository = FakeRepository()
    monkeypatch.setattr(
        app_module,
        "get_flood_repository",
        lambda include_prediction=False: repository,
    )
    monkeypatch.setattr(app_module, "load_active_location_reports", lambda: [])
    return repository


def test_haversine_distance_nearby_points():
    distance = haversine_m(36.1289, 128.3309, 36.1195, 128.3447)
    assert 1500 < distance < 1800


def test_hazard_detection_and_detour_generation(fake_flood_repository):
    start = {"lat": 36.1289, "lng": 128.3309}
    end = {"lat": 36.11953, "lng": 128.34474}

    hazards = hazards_for_segment(start, end)
    assert hazards

    waypoints = detour_waypoints(start, end, hazards)
    assert waypoints
    assert {"lat", "lng"} <= set(waypoints[0])
    assert hazards_for_path([start, *waypoints, end]) == []


def test_car_candidates_keep_boundary_and_corridor_routes(fake_flood_repository):
    start = {"lat": 36.1289, "lng": 128.3309}
    end = {"lat": 36.11953, "lng": 128.34474}
    hazards = hazards_for_segment(start, end)

    candidates = detour_waypoint_candidates(
        start,
        end,
        hazards,
        limit=8,
        mode="car",
        repository=fake_flood_repository,
    )
    boundary_count = sum(
        len({round(point["lat"], 6) for point in candidate}) == 1
        or len({round(point["lng"], 6) for point in candidate}) == 1
        for candidate in candidates
    )

    assert len(candidates) == 8
    assert boundary_count == 4


def test_history_and_optional_prediction_outputs_are_exposed():
    repository = get_flood_repository()
    combined_repository = get_flood_repository(include_prediction=True)
    fifty_year_repository = FloodRepository(
        include_prediction=True,
        include_history=False,
    )
    display = load_flood_display()

    assert set(repository.source_ids) == {"history"}
    assert set(combined_repository.source_ids) == {"history", "river", "urban"}
    assert set(fifty_year_repository.source_ids) == {"river", "urban"}
    assert len(combined_repository.geometries) > len(repository.geometries)
    assert display["type"] == "FeatureCollection"
    assert {feature["properties"]["id"] for feature in display["features"]} == {
        "history",
        "river",
        "urban",
    }
    assert display["features"][0]["properties"]["datasetType"] == "history"
    assert display["features"][1]["properties"]["datasetType"] == "scenario"


def test_flood_exposure_distance_counts_each_traversed_segment():
    repository = get_flood_repository(include_prediction=True)
    start = {"lat": 36.1289, "lng": 128.3309}
    end = {"lat": 36.11953, "lng": 128.34474}

    single = repository.risk_exposure_m([[start, end]])
    repeated = repository.risk_exposure_m([[start, end], [end, start]])

    assert single > 0
    assert repeated == pytest.approx(single * 2)


def test_recent_flood_trace_filter_uses_gumi_and_requested_years():
    record = {
        "STDG_SGG_CD": "47190",
        "FLDN_YR": "2022",
        "GEOM": "POLYGON EMPTY",
    }

    assert record_matches(record, "47190", 2021, 2025)
    assert not record_matches({**record, "STDG_SGG_CD": "47130"}, "47190", 2021, 2025)
    assert not record_matches({**record, "FLDN_YR": "2018"}, "47190", 2021, 2025)


def test_shelter_filter_keeps_only_gumi_temporary_housing_records():
    record = {
        "ARCD": "4719000000",
        "BDONG_CD": "4719010100",
        "LA": "36.12",
        "LO": "128.34",
    }

    assert is_gumi_temporary_housing(record)
    assert not is_gumi_temporary_housing(
        {**record, "ARCD": "4713000000", "BDONG_CD": "4713010100"}
    )
    assert not is_gumi_temporary_housing({**record, "LA": None})


def test_official_shelter_data_contains_only_verified_temporary_housing():
    shelters = app_module.load_shelters()

    assert len({shelter["id"] for shelter in shelters}) == len(shelters)
    assert all(shelter["verified"] is True for shelter in shelters)
    assert all(shelter["type"] == "이재민 임시주거시설" for shelter in shelters)
    assert all(
        shelter["source"] == "safetydata_DSSP-IF-10945"
        for shelter in shelters
    )
    assert all(35.7 < shelter["lat"] < 36.5 for shelter in shelters)
    assert all(127.7 < shelter["lng"] < 129.0 for shelter in shelters)


def test_location_report_is_saved_and_returned_to_admin(report_db):
    client = app_module.app.test_client()
    response = client.post(
        "/api/reports",
        json={
            "lat": 36.11856,
            "lng": 128.36556,
            "accuracyM": 14.2,
            "note": "도로 일부 침수",
        },
    )
    payload = response.get_json()

    assert response.status_code == 201
    assert report_db.exists()
    assert payload["report"]["id"] == 1
    assert payload["report"]["accuracyM"] == 14.2

    list_response = client.get("/api/admin/reports")
    list_payload = list_response.get_json()
    assert list_response.status_code == 200
    assert list_payload["count"] == 1
    assert list_payload["reports"][0]["note"] == "도로 일부 침수"


def test_location_report_rejects_coordinates_outside_gumi_area(report_db):
    response = app_module.app.test_client().post(
        "/api/reports",
        json={"lat": 37.5665, "lng": 126.9780},
    )

    assert response.status_code == 400
    assert "구미시" in response.get_json()["error"]
    assert not report_db.exists()


def test_admin_reports_page_is_available():
    response = app_module.app.test_client().get("/admin/reports")
    html = response.get_data(as_text=True)

    assert response.status_code == 200
    assert "신고 좌표 관리" in html
    assert 'id="startFloodPlacement"' in html
    assert 'id="adminFloodDialog"' in html
    assert 'id="dongStatsTab"' in html
    assert 'id="dongStatsList"' in html
    assert 'id="openHistoryStatistics"' in html
    assert 'id="historyStatisticsDialog"' in html
    assert 'id="historyStatisticsChart"' in html


def test_user_page_defaults_to_current_flood_only():
    html = app_module.app.test_client().get("/").get_data(as_text=True)

    assert 'id="includeLiveFlood"' not in html
    assert 'id="includeRiskAreas" type="checkbox"' in html
    assert 'id="includeRiskAreas" type="checkbox" checked' not in html
    assert 'id="includeFiftyYear" type="checkbox"' in html
    assert 'id="includeFiftyYear" type="checkbox" checked' not in html


def test_risk_area_request_uses_history_without_prediction(
    monkeypatch, route_shelter
):
    repository_calls = []

    def capture_repository(include_prediction, include_live_flood, include_history):
        repository_calls.append(
            (include_prediction, include_live_flood, include_history)
        )
        return object()

    monkeypatch.setattr(
        app_module,
        "flood_repository_for_request",
        capture_repository,
    )

    client = app_module.app.test_client()
    history_response = client.get(
        "/api/shelters?includeRiskAreas=true&includePrediction=true"
    )
    fifty_year_response = client.get(
        "/api/shelters?includeFiftyYear=true&includeLiveFlood=false"
    )

    assert history_response.status_code == 200
    assert fifty_year_response.status_code == 200
    assert repository_calls == [
        (False, True, True),
        (True, True, False),
    ]


def test_report_approval_and_resolution_control_live_flood_zone(report_db):
    client = app_module.app.test_client()
    created = client.post(
        "/api/reports",
        json={"lat": 36.11856, "lng": 128.36556, "accuracyM": 9},
    ).get_json()["report"]

    assert created["status"] == "pending"
    assert client.get("/api/live-flood-zones").get_json()["features"] == []

    approved_response = client.post(
        f"/api/admin/reports/{created['id']}/approve"
    )
    approved = approved_response.get_json()["report"]
    live_payload = client.get("/api/live-flood-zones").get_json()

    assert approved_response.status_code == 200
    assert approved["status"] == "active"
    assert approved["approvedAt"]
    assert len(live_payload["features"]) == 1
    assert live_payload["features"][0]["properties"]["datasetType"] == "live"
    assert live_payload["features"][0]["properties"]["radiusM"] == 40

    resolved_response = client.post(
        f"/api/admin/reports/{created['id']}/resolve"
    )
    assert resolved_response.status_code == 200
    assert resolved_response.get_json()["report"]["status"] == "resolved"
    assert client.get("/api/live-flood-zones").get_json()["features"] == []


def test_admin_map_point_is_immediately_added_as_active_flood(report_db):
    client = app_module.app.test_client()
    response = client.post(
        "/api/admin/flood-zones",
        json={
            "lat": 36.11856,
            "lng": 128.36556,
            "note": "관리자 현장 확인",
        },
    )
    report = response.get_json()["report"]
    live_payload = client.get("/api/live-flood-zones").get_json()

    assert response.status_code == 201
    assert report["status"] == "active"
    assert report["source"] == "admin"
    assert report["isAdminCreated"] is True
    assert report["approvedAt"]
    assert len(live_payload["features"]) == 1
    properties = live_payload["features"][0]["properties"]
    assert properties["reportId"] == report["id"]
    assert properties["isAdminCreated"] is True
    assert properties["radiusM"] == 40
    assert properties["name"].startswith("관리자 등록 침수")
    export_response = client.get("/api/admin/reports/export.xlsx")
    workbook = load_workbook(BytesIO(export_response.data))
    assert workbook["침수 신고 이력"].cell(row=2, column=13).value == (
        "관리자 직접 추가"
    )


def test_admin_map_point_rejects_coordinates_outside_gumi(report_db):
    response = app_module.app.test_client().post(
        "/api/admin/flood-zones",
        json={"lat": 37.5665, "lng": 126.9780},
    )

    assert response.status_code == 400
    assert "구미시" in response.get_json()["error"]
    assert not report_db.exists()


def test_gumi_admin_districts_cover_all_demo_flood_points():
    repository = get_admin_district_repository()
    reports = [
        {
            "id": index,
            "lat": lat,
            "lng": lng,
            "isSample": True,
            "isAdminCreated": False,
        }
        for index, (lat, lng) in enumerate(app_module.DEMO_FLOOD_POINTS, start=1)
    ]
    payload = repository.statistics(reports)

    assert payload["metadata"]["featureCount"] == 25
    assert payload["metadata"]["activeCount"] == 15
    assert payload["metadata"]["unmatchedCount"] == 0
    assert sum(
        feature["properties"]["activeCount"]
        for feature in payload["features"]
    ) == 15
    assert all(
        feature["properties"]["fullName"].startswith("경상북도 구미시 ")
        for feature in payload["features"]
    )


def test_admin_dong_statistics_follow_active_flood_status(report_db):
    client = app_module.app.test_client()
    report = client.post(
        "/api/admin/flood-zones",
        json={"lat": 36.11856, "lng": 128.36556},
    ).get_json()["report"]

    active_payload = client.get(
        "/api/admin/flood-statistics/dongs"
    ).get_json()
    active_features = [
        feature
        for feature in active_payload["features"]
        if feature["properties"]["activeCount"]
    ]

    assert active_payload["metadata"]["activeCount"] == 1
    assert active_payload["metadata"]["affectedDongCount"] == 1
    assert active_payload["metadata"]["unmatchedCount"] == 0
    assert active_features[0]["properties"]["adminCount"] == 1
    assert active_features[0]["properties"]["reportIds"] == [report["id"]]

    client.post(f"/api/admin/reports/{report['id']}/resolve")
    resolved_payload = client.get(
        "/api/admin/flood-statistics/dongs"
    ).get_json()
    assert resolved_payload["metadata"]["activeCount"] == 0
    assert resolved_payload["metadata"]["affectedDongCount"] == 0


def test_admin_dong_history_statistics_use_confirmed_sql_records(report_db):
    client = app_module.app.test_client()
    active_report = client.post(
        "/api/admin/flood-zones",
        json={"lat": 36.11856, "lng": 128.36556},
    ).get_json()["report"]
    resolved_report = client.post(
        "/api/admin/flood-zones",
        json={"lat": 36.11860, "lng": 128.36560},
    ).get_json()["report"]
    client.post(f"/api/admin/reports/{resolved_report['id']}/resolve")
    client.post(
        "/api/reports",
        json={"lat": 36.11864, "lng": 128.36564},
    )

    response = client.get("/api/admin/flood-statistics/history/dongs")
    payload = response.get_json()
    districts_with_history = [
        district for district in payload["districts"] if district["totalCount"]
    ]

    assert response.status_code == 200
    assert payload["metadata"]["totalCount"] == 2
    assert payload["metadata"]["activeCount"] == 1
    assert payload["metadata"]["resolvedCount"] == 1
    assert payload["metadata"]["affectedDongCount"] == 1
    assert payload["metadata"]["unmatchedCount"] == 0
    assert payload["metadata"]["topDongName"]
    assert len(payload["districts"]) == 25
    assert len(districts_with_history) == 1
    assert districts_with_history[0]["activeCount"] == 1
    assert districts_with_history[0]["resolvedCount"] == 1
    assert active_report["id"] != resolved_report["id"]


def test_live_flood_geometry_uses_approximately_40_meter_radius():
    lat = 36.11856
    lng = 128.36556
    geometry = live_flood_geometry(lat, lng)
    edge_lng, edge_lat = geometry.exterior.coords[0]

    radius = haversine_m(lat, lng, edge_lat, edge_lng)
    assert radius == pytest.approx(40, abs=0.5)


def test_demo_flood_seed_adds_15_idempotent_labeled_points(report_db):
    first = app_module.seed_demo_flood_reports()
    second = app_module.seed_demo_flood_reports()
    reports = app_module.load_active_location_reports()
    demo_reports = [report for report in reports if report["isSample"]]
    history_repository = get_flood_repository()
    live_features = app_module.live_flood_feature_collection(demo_reports)

    assert first == {"inserted": 15, "total": 15}
    assert second == {"inserted": 0, "total": 15}
    assert len(demo_reports) == 15
    assert len({report["sampleKey"] for report in demo_reports}) == 15
    assert all(report["status"] == "active" for report in demo_reports)
    assert all(
        history_repository.contains_point(
            {"lat": report["lat"], "lng": report["lng"]}
        )
        for report in demo_reports
    )
    assert all(
        feature["properties"]["isSample"] is True
        and feature["properties"]["riskLevel"] == "데모 침수구역"
        for feature in live_features["features"]
    )


def test_approved_live_flood_zone_is_used_for_route_hazard_checks(report_db):
    report = app_module.save_location_report(
        lat=36.11856,
        lng=128.36556,
        accuracy_m=8,
        note="경로 교차 테스트",
    )
    app_module.update_location_report_status(report["id"], "approve")
    repository = app_module.flood_repository_for_request(
        include_prediction=False,
        include_live_flood=True,
    )

    hazards = repository.hazards_for_paths(
        [[
            {"lat": 36.11856, "lng": 128.36456},
            {"lat": 36.11856, "lng": 128.36656},
        ]]
    )

    assert any(
        hazard["sourceId"] == "live" and hazard["reportId"] == report["id"]
        for hazard in hazards
    )


def test_current_flood_only_repository_excludes_static_risk_areas(report_db):
    report = app_module.save_location_report(
        lat=36.11856,
        lng=128.36556,
        accuracy_m=8,
        note="현재 침수 전용 테스트",
    )
    app_module.update_location_report_status(report["id"], "approve")

    repository = app_module.flood_repository_for_request(
        include_prediction=False,
        include_live_flood=True,
        include_history=False,
    )

    assert set(repository.source_ids) == {"live"}


def test_excel_export_contains_confirmed_history_and_excludes_pending(report_db):
    pending = app_module.save_location_report(
        lat=36.11001,
        lng=128.35001,
        accuracy_m=12,
        note="승인 대기",
    )
    confirmed = app_module.save_location_report(
        lat=36.12002,
        lng=128.36002,
        accuracy_m=7,
        note="확정 침수",
    )
    app_module.update_location_report_status(confirmed["id"], "approve")
    app_module.update_location_report_status(confirmed["id"], "resolve")

    response = app_module.app.test_client().get(
        "/api/admin/reports/export.xlsx"
    )
    workbook = load_workbook(BytesIO(response.data))
    sheet = workbook["침수 신고 이력"]
    rows = list(sheet.iter_rows(values_only=True))

    assert response.status_code == 200
    assert response.mimetype == (
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert rows[0][:7] == (
        "신고번호",
        "상태",
        "발생연도",
        "발생일자",
        "발생시간",
        "위도",
        "경도",
    )
    assert [row[0] for row in rows[1:]] == [confirmed["id"]]
    assert rows[1][1] == "침수 해제"
    assert rows[1][2] == app_module._report_datetime_in_kst(
        confirmed["createdAt"]
    ).year
    assert pending["id"] not in [row[0] for row in rows[1:]]


def test_tmap_route_excludes_unindexed_waypoint_connector_lines():
    payload = {
        "features": [
            {
                "geometry": {
                    "type": "LineString",
                    "coordinates": [[128.01, 36.0], [128.02, 36.0]],
                },
                "properties": {"index": 3, "lineIndex": 1},
            },
            {
                "geometry": {
                    "type": "LineString",
                    "coordinates": [[128.0, 36.0], [128.01, 36.0]],
                },
                "properties": {"index": 1, "lineIndex": 0},
            },
            {
                "geometry": {
                    "type": "LineString",
                    "coordinates": [[129.0, 37.0], [128.01, 36.0]],
                },
                "properties": {},
            },
        ]
    }

    route = _extract_tmap_route(payload)

    assert len(route["pathSegments"]) == 2
    assert route["path"] == [
        {"lng": 128.0, "lat": 36.0},
        {"lng": 128.01, "lat": 36.0},
        {"lng": 128.02, "lat": 36.0},
    ]


@pytest.mark.parametrize("mode", ["pedestrian", "car"])
def test_route_api_rechecks_tmap_path_and_retries_with_detour(
    monkeypatch, fake_flood_repository, route_shelter, mode
):
    calls = []

    def fake_tmap_route(start, end, end_name, pass_points, mode):
        calls.append(pass_points)
        path = [start, *pass_points, end]
        return (
            {
                "source": "tmap",
                "summary": {
                    "totalDistance": round(
                        sum(
                            haversine_m(
                                a["lat"], a["lng"], b["lat"], b["lng"]
                            )
                            for a, b in zip(path, path[1:])
                        )
                    ),
                    "totalTime": 600,
                },
                "path": path,
                "steps": [],
            },
            None,
        )

    monkeypatch.setattr(app_module, "request_tmap_route", fake_tmap_route)
    client = app_module.app.test_client()
    response = client.post(
        "/api/route",
        json={
            "start": {"lat": 36.1289, "lng": 128.3309},
            "shelterId": route_shelter["id"],
            "mode": mode,
            "avoidFlood": True,
            "includeRiskAreas": True,
        },
    )
    payload = response.get_json()

    assert response.status_code == 200
    assert calls[0] == []
    assert len(calls) >= 2
    assert payload["avoidanceStatus"] == "clear"
    assert payload["riskZones"] == []
    assert payload["avoidedRiskZones"][0]["id"] == "flood-demo-001"
    assert len(payload["passPoints"]) == 3


def test_route_api_prefers_less_flood_exposure_over_fewer_zone_ids(
    monkeypatch, route_shelter
):
    zones = [
        {
            "id": f"flood-{index}",
            "sourceId": "test",
            "name": f"테스트 침수구역 {index}",
            "riskLevel": "테스트",
            "note": "테스트용 도형",
            "detour_margin_deg": 0.0015,
            "bounds": [128.336, 36.122, 128.340, 36.127],
        }
        for index in range(3)
    ]

    class ExposureRepository:
        @staticmethod
        def _is_detour(paths):
            return any(len(path) > 2 for path in paths)

        @classmethod
        def hazards_for_paths(cls, paths):
            return zones if cls._is_detour(paths) else zones[:1]

        @staticmethod
        def hazards_for_point(point):
            return []

        @staticmethod
        def contains_point(point):
            return False

        @staticmethod
        def segment_intersects_zone(start, end, zone):
            return True

        @classmethod
        def risk_exposure_m(cls, paths):
            return 80 if cls._is_detour(paths) else 800

    repository = ExposureRepository()
    monkeypatch.setattr(
        app_module,
        "get_flood_repository",
        lambda include_prediction=False: repository,
    )
    monkeypatch.setattr(app_module, "load_active_location_reports", lambda: [])

    def fake_tmap_route(start, end, end_name, pass_points, mode):
        path = [start, *pass_points, end]
        return (
            {
                "source": "tmap",
                "summary": {"totalDistance": 2000, "totalTime": 1200},
                "path": path,
                "pathSegments": [path],
                "steps": [],
            },
            None,
        )

    monkeypatch.setattr(app_module, "request_tmap_route", fake_tmap_route)
    response = app_module.app.test_client().post(
        "/api/route",
        json={
            "start": {"lat": 36.1289, "lng": 128.3309},
            "shelterId": route_shelter["id"],
            "mode": "pedestrian",
            "avoidFlood": True,
            "includeRiskAreas": True,
        },
    )
    payload = response.get_json()

    assert response.status_code == 200
    assert payload["passPoints"]
    assert len(payload["riskZones"]) == 3
    assert payload["initialRiskDistanceM"] == 800
    assert payload["riskDistanceM"] == 80
    assert payload["riskReductionPercent"] == 90

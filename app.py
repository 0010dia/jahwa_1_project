from __future__ import annotations

import json
import math
import os
import sqlite3
from datetime import UTC, datetime, timedelta, timezone
from functools import lru_cache
from io import BytesIO
from itertools import product
from pathlib import Path
from typing import Any
from urllib.parse import quote

import requests
from flask import Flask, jsonify, render_template, request, send_file
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from flood_data import (
    FloodRepository,
    get_flood_repository,
    live_flood_feature_collection,
    load_flood_display,
)

try:
    from dotenv import load_dotenv
except ImportError:  # pragma: no cover - optional local convenience
    def load_dotenv() -> None:
        return None


BASE_DIR = Path(__file__).resolve().parent
DATA_DIR = BASE_DIR / "data"
EARTH_RADIUS_M = 6_371_000
KST = timezone(timedelta(hours=9))

load_dotenv()

app = Flask(__name__)
app.config["JSON_AS_ASCII"] = False
REPORTS_DB_PATH = Path(
    os.getenv("REPORTS_DB_PATH", str(DATA_DIR / "reports.sqlite3"))
)
REPORT_BOUNDS = {
    "minLat": 35.7,
    "maxLat": 36.5,
    "minLng": 127.7,
    "maxLng": 129.0,
}
REPORT_NOTE_MAX_LENGTH = 300
DEMO_FLOOD_POINTS = [
    (36.20081192, 128.28469764),
    (36.20228105, 128.28386727),
    (36.20107244, 128.28716587),
    (36.19994068, 128.27414245),
    (36.20136001, 128.26274945),
    (36.20082690, 128.27290516),
    (36.20406418, 128.23988772),
    (36.20431277, 128.24206044),
    (36.20372029, 128.23565082),
    (36.10325805, 128.37047393),
    (36.13662005, 128.32969330),
    (36.13061771, 128.33735393),
    (36.10736043, 128.41838483),
    (36.10714220, 128.42120536),
    (36.13555607, 128.33405949),
]


def _load_json(filename: str) -> list[dict[str, Any]]:
    with (DATA_DIR / filename).open(encoding="utf-8") as file:
        return json.load(file)


@lru_cache(maxsize=1)
def load_shelters() -> list[dict[str, Any]]:
    return _load_json("shelters.gumi.json")


def _report_connection() -> sqlite3.Connection:
    REPORTS_DB_PATH.parent.mkdir(parents=True, exist_ok=True)
    connection = sqlite3.connect(REPORTS_DB_PATH, timeout=5)
    connection.row_factory = sqlite3.Row
    connection.execute(
        """
        CREATE TABLE IF NOT EXISTS location_reports (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            lat REAL NOT NULL,
            lng REAL NOT NULL,
            accuracy_m REAL,
            note TEXT NOT NULL DEFAULT '',
            created_at TEXT NOT NULL,
            status TEXT NOT NULL DEFAULT 'pending',
            approved_at TEXT,
            resolved_at TEXT,
            source TEXT NOT NULL DEFAULT 'user',
            sample_key TEXT
        )
        """
    )
    columns = {
        str(row["name"])
        for row in connection.execute("PRAGMA table_info(location_reports)")
    }
    migrations = {
        "status": "ALTER TABLE location_reports ADD COLUMN status TEXT NOT NULL DEFAULT 'pending'",
        "approved_at": "ALTER TABLE location_reports ADD COLUMN approved_at TEXT",
        "resolved_at": "ALTER TABLE location_reports ADD COLUMN resolved_at TEXT",
        "source": "ALTER TABLE location_reports ADD COLUMN source TEXT NOT NULL DEFAULT 'user'",
        "sample_key": "ALTER TABLE location_reports ADD COLUMN sample_key TEXT",
    }
    for column, statement in migrations.items():
        if column not in columns:
            connection.execute(statement)
    return connection


def _serialize_report(row: sqlite3.Row) -> dict[str, Any]:
    return {
        "id": int(row["id"]),
        "lat": float(row["lat"]),
        "lng": float(row["lng"]),
        "accuracyM": (
            float(row["accuracy_m"])
            if row["accuracy_m"] is not None
            else None
        ),
        "note": str(row["note"] or ""),
        "createdAt": str(row["created_at"]),
        "status": str(row["status"]),
        "approvedAt": (
            str(row["approved_at"]) if row["approved_at"] is not None else None
        ),
        "resolvedAt": (
            str(row["resolved_at"]) if row["resolved_at"] is not None else None
        ),
        "source": str(row["source"] or "user"),
        "sampleKey": (
            str(row["sample_key"]) if row["sample_key"] is not None else None
        ),
        "isSample": str(row["source"] or "user") == "demo",
        "isAdminCreated": str(row["source"] or "user") == "admin",
    }


def save_location_report(
    lat: float,
    lng: float,
    accuracy_m: float | None,
    note: str,
    source: str = "user",
    sample_key: str | None = None,
) -> dict[str, Any]:
    created_at = datetime.now(UTC).isoformat(timespec="seconds")
    with _report_connection() as connection:
        cursor = connection.execute(
            """
            INSERT INTO location_reports (
                lat, lng, accuracy_m, note, created_at, source, sample_key
            )
            VALUES (?, ?, ?, ?, ?, ?, ?)
            """,
            (lat, lng, accuracy_m, note, created_at, source, sample_key),
        )
        report_id = int(cursor.lastrowid)
    return {
        "id": report_id,
        "lat": lat,
        "lng": lng,
        "accuracyM": accuracy_m,
        "note": note,
        "createdAt": created_at,
        "status": "pending",
        "approvedAt": None,
        "resolvedAt": None,
        "source": source,
        "sampleKey": sample_key,
        "isSample": source == "demo",
        "isAdminCreated": source == "admin",
    }


def seed_demo_flood_reports() -> dict[str, int]:
    now = datetime.now(UTC).isoformat(timespec="seconds")
    inserted = 0
    with _report_connection() as connection:
        existing_keys = {
            str(row["sample_key"])
            for row in connection.execute(
                """
                SELECT sample_key FROM location_reports
                WHERE source = 'demo' AND sample_key IS NOT NULL
                """
            ).fetchall()
        }
        for index, (lat, lng) in enumerate(DEMO_FLOOD_POINTS, start=1):
            sample_key = f"history-demo-{index:02d}"
            if sample_key in existing_keys:
                continue
            connection.execute(
                """
                INSERT INTO location_reports (
                    lat, lng, accuracy_m, note, created_at, status,
                    approved_at, resolved_at, source, sample_key
                )
                VALUES (?, ?, NULL, ?, ?, 'active', ?, NULL, 'demo', ?)
                """,
                (
                    lat,
                    lng,
                    f"실제 침수흔적 내부 추출 데모 지점 {index:02d}",
                    now,
                    now,
                    sample_key,
                ),
            )
            inserted += 1
    return {"inserted": inserted, "total": len(DEMO_FLOOD_POINTS)}


if os.getenv("AUTO_SEED_DEMO_FLOODS", "true").strip().lower() in {
    "1",
    "true",
    "yes",
    "on",
}:
    seed_demo_flood_reports()


def load_location_reports(limit: int = 500) -> list[dict[str, Any]]:
    with _report_connection() as connection:
        rows = connection.execute(
            """
            SELECT id, lat, lng, accuracy_m, note, created_at,
                   status, approved_at, resolved_at, source, sample_key
            FROM location_reports
            ORDER BY id DESC
            LIMIT ?
            """,
            (limit,),
        ).fetchall()
    return [_serialize_report(row) for row in rows]


def load_active_location_reports() -> list[dict[str, Any]]:
    with _report_connection() as connection:
        rows = connection.execute(
            """
            SELECT id, lat, lng, accuracy_m, note, created_at,
                   status, approved_at, resolved_at, source, sample_key
            FROM location_reports
            WHERE status = 'active'
            ORDER BY approved_at DESC, id DESC
            """
        ).fetchall()
    return [_serialize_report(row) for row in rows]


def load_confirmed_flood_reports() -> list[dict[str, Any]]:
    with _report_connection() as connection:
        rows = connection.execute(
            """
            SELECT id, lat, lng, accuracy_m, note, created_at,
                   status, approved_at, resolved_at, source, sample_key
            FROM location_reports
            WHERE status IN ('active', 'resolved')
            ORDER BY created_at DESC, id DESC
            """
        ).fetchall()
    return [_serialize_report(row) for row in rows]


def _report_datetime_in_kst(value: str | None) -> datetime | None:
    if not value:
        return None
    parsed = datetime.fromisoformat(value)
    if parsed.tzinfo is None:
        parsed = parsed.replace(tzinfo=UTC)
    return parsed.astimezone(KST).replace(tzinfo=None)


def _excel_safe_text(value: str) -> str:
    if value.startswith(("=", "+", "-", "@")):
        return f"'{value}"
    return value


def build_flood_history_workbook(reports: list[dict[str, Any]]) -> BytesIO:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "침수 신고 이력"
    headers = [
        "신고번호",
        "상태",
        "발생연도",
        "발생일자",
        "발생시간",
        "위도",
        "경도",
        "GPS 정확도(m)",
        "현장 메모",
        "신고일시",
        "승인일시",
        "해제일시",
        "데이터 구분",
    ]
    sheet.append(headers)
    status_labels = {"active": "현재 침수", "resolved": "침수 해제"}

    for report in reports:
        reported_at = _report_datetime_in_kst(report["createdAt"])
        approved_at = _report_datetime_in_kst(report.get("approvedAt"))
        resolved_at = _report_datetime_in_kst(report.get("resolvedAt"))
        sheet.append(
            [
                report["id"],
                status_labels.get(report["status"], report["status"]),
                reported_at.year if reported_at else None,
                reported_at.date() if reported_at else None,
                reported_at.time() if reported_at else None,
                report["lat"],
                report["lng"],
                report["accuracyM"],
                _excel_safe_text(report["note"]),
                reported_at,
                approved_at,
                resolved_at,
                (
                    "데모"
                    if report.get("isSample")
                    else "관리자 직접 추가"
                    if report.get("isAdminCreated")
                    else "사용자 신고"
                ),
            ]
        )

    header_fill = PatternFill("solid", fgColor="155E75")
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center")

    for row in sheet.iter_rows(min_row=2):
        row[3].number_format = "yyyy-mm-dd"
        row[4].number_format = "hh:mm:ss"
        row[5].number_format = "0.000000"
        row[6].number_format = "0.000000"
        for cell in row[9:12]:
            cell.number_format = "yyyy-mm-dd hh:mm:ss"

    widths = [10, 12, 10, 13, 12, 13, 13, 15, 34, 21, 21, 21, 14]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[chr(64 + index)].width = width
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions

    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output


def update_location_report_status(
    report_id: int, action: str
) -> dict[str, Any] | None:
    now = datetime.now(UTC).isoformat(timespec="seconds")
    with _report_connection() as connection:
        existing = connection.execute(
            "SELECT status FROM location_reports WHERE id = ?", (report_id,)
        ).fetchone()
        if existing is None:
            return None
        if action == "resolve" and existing["status"] != "active":
            raise ValueError("활성화된 침수 신고만 해제할 수 있습니다.")
        if action == "approve":
            connection.execute(
                """
                UPDATE location_reports
                SET status = 'active', approved_at = ?, resolved_at = NULL
                WHERE id = ?
                """,
                (now, report_id),
            )
        elif action == "resolve":
            connection.execute(
                """
                UPDATE location_reports
                SET status = 'resolved', resolved_at = ?
                WHERE id = ?
                """,
                (now, report_id),
            )
        else:
            raise ValueError("지원하지 않는 신고 상태 변경입니다.")
        row = connection.execute(
            """
            SELECT id, lat, lng, accuracy_m, note, created_at,
                   status, approved_at, resolved_at, source, sample_key
            FROM location_reports WHERE id = ?
            """,
            (report_id,),
        ).fetchone()
    return _serialize_report(row)


def flood_repository_for_request(
    include_prediction: bool,
    include_live_flood: bool,
    include_history: bool = True,
) -> Any:
    if include_history and not include_live_flood:
        return get_flood_repository(include_prediction)
    live_reports = load_active_location_reports() if include_live_flood else []
    if include_history and not live_reports:
        return get_flood_repository(include_prediction)
    return FloodRepository(
        include_prediction=include_prediction,
        live_reports=live_reports,
        include_history=include_history,
    )


def haversine_m(lat1: float, lng1: float, lat2: float, lng2: float) -> float:
    lat1_rad, lat2_rad = math.radians(lat1), math.radians(lat2)
    dlat = math.radians(lat2 - lat1)
    dlng = math.radians(lng2 - lng1)
    a = (
        math.sin(dlat / 2) ** 2
        + math.cos(lat1_rad) * math.cos(lat2_rad) * math.sin(dlng / 2) ** 2
    )
    return EARTH_RADIUS_M * 2 * math.atan2(math.sqrt(a), math.sqrt(1 - a))


def _xy(point: dict[str, float], origin_lat: float) -> tuple[float, float]:
    x = math.radians(point["lng"]) * EARTH_RADIUS_M * math.cos(math.radians(origin_lat))
    y = math.radians(point["lat"]) * EARTH_RADIUS_M
    return x, y


def _orientation(
    a: tuple[float, float], b: tuple[float, float], c: tuple[float, float]
) -> float:
    return (b[1] - a[1]) * (c[0] - b[0]) - (b[0] - a[0]) * (c[1] - b[1])


def _on_segment(
    a: tuple[float, float], b: tuple[float, float], c: tuple[float, float]
) -> bool:
    return (
        min(a[0], c[0]) <= b[0] <= max(a[0], c[0])
        and min(a[1], c[1]) <= b[1] <= max(a[1], c[1])
    )


def _segments_intersect(
    p1: tuple[float, float],
    q1: tuple[float, float],
    p2: tuple[float, float],
    q2: tuple[float, float],
) -> bool:
    o1 = _orientation(p1, q1, p2)
    o2 = _orientation(p1, q1, q2)
    o3 = _orientation(p2, q2, p1)
    o4 = _orientation(p2, q2, q1)
    epsilon = 1e-9

    if o1 * o2 < 0 and o3 * o4 < 0:
        return True
    if abs(o1) < epsilon and _on_segment(p1, p2, q1):
        return True
    if abs(o2) < epsilon and _on_segment(p1, q2, q1):
        return True
    if abs(o3) < epsilon and _on_segment(p2, p1, q2):
        return True
    if abs(o4) < epsilon and _on_segment(p2, q1, q2):
        return True
    return False


def point_in_polygon(point: dict[str, float], polygon: list[dict[str, float]]) -> bool:
    inside = False
    x, y = point["lng"], point["lat"]
    previous = polygon[-1]

    for current in polygon:
        xi, yi = current["lng"], current["lat"]
        xj, yj = previous["lng"], previous["lat"]
        intersects = (yi > y) != (yj > y) and x < (xj - xi) * (y - yi) / (
            yj - yi + 1e-12
        ) + xi
        if intersects:
            inside = not inside
        previous = current

    return inside


def segment_intersects_polygon(
    start: dict[str, float], end: dict[str, float], polygon: list[dict[str, float]]
) -> bool:
    if point_in_polygon(start, polygon) or point_in_polygon(end, polygon):
        return True

    origin_lat = (start["lat"] + end["lat"]) / 2
    line_start = _xy(start, origin_lat)
    line_end = _xy(end, origin_lat)
    previous = polygon[-1]

    for current in polygon:
        edge_start = _xy(previous, origin_lat)
        edge_end = _xy(current, origin_lat)
        if _segments_intersect(line_start, line_end, edge_start, edge_end):
            return True
        previous = current

    return False


def hazards_for_segment(
    start: dict[str, float],
    end: dict[str, float],
    include_prediction: bool = False,
) -> list[dict[str, Any]]:
    return get_flood_repository(include_prediction).hazards_for_paths([[start, end]])


def hazards_for_path(
    path: list[dict[str, float]], include_prediction: bool = False
) -> list[dict[str, Any]]:
    return get_flood_repository(include_prediction).hazards_for_paths([path])


def hazards_for_route(
    route: dict[str, Any], repository: Any | None = None
) -> list[dict[str, Any]]:
    segments = route.get("pathSegments") or [route.get("path", [])]
    repository = repository or get_flood_repository()
    return repository.hazards_for_paths(segments)


def _route_risk_distance_m(route: dict[str, Any], repository: Any) -> float:
    calculator = getattr(repository, "risk_exposure_m", None)
    if not callable(calculator):
        return 0.0
    segments = route.get("pathSegments") or [route.get("path", [])]
    return float(calculator(segments))


def _route_projection_fraction(
    point: dict[str, float], start: dict[str, float], end: dict[str, float]
) -> float:
    origin_lat = (start["lat"] + end["lat"]) / 2
    px, py = _xy(point, origin_lat)
    sx, sy = _xy(start, origin_lat)
    ex, ey = _xy(end, origin_lat)
    dx, dy = ex - sx, ey - sy
    denominator = dx * dx + dy * dy
    if denominator == 0:
        return 0
    return max(0, min(1, ((px - sx) * dx + (py - sy) * dy) / denominator))


def _zone_center(zone: dict[str, Any]) -> dict[str, float]:
    min_lng, min_lat, max_lng, max_lat = zone["bounds"]
    return {
        "lat": (min_lat + max_lat) / 2,
        "lng": (min_lng + max_lng) / 2,
    }


def _path_distance(path: list[dict[str, float]]) -> float:
    return sum(
        haversine_m(
            start["lat"], start["lng"], end["lat"], end["lng"]
        )
        for start, end in zip(path, path[1:])
    )


def _zone_detour_options(
    start: dict[str, float],
    end: dict[str, float],
    zone: dict[str, Any],
    mode: str = "pedestrian",
    repository: Any | None = None,
) -> list[list[dict[str, float]]]:
    repository = repository or get_flood_repository()
    min_lng, min_lat, max_lng, max_lat = zone["bounds"]
    base_margin = float(zone.get("detour_margin_deg", 0.0045))
    margins = [base_margin]
    if mode == "car":
        margins = sorted(
            {
                max(base_margin * 1.5, 0.006),
                max(base_margin * 2, 0.008),
            }
        )

    center_lat = (min_lat + max_lat) / 2
    center_lng = (min_lng + max_lng) / 2
    options = []
    for margin in margins:
        south_west = {"lat": min_lat - margin, "lng": min_lng - margin}
        south_east = {"lat": min_lat - margin, "lng": max_lng + margin}
        north_west = {"lat": max_lat + margin, "lng": min_lng - margin}
        north_east = {"lat": max_lat + margin, "lng": max_lng + margin}
        options.extend(
            [
                [
                    south_west,
                    {"lat": min_lat - margin, "lng": center_lng},
                    south_east,
                ],
                [
                    north_west,
                    {"lat": max_lat + margin, "lng": center_lng},
                    north_east,
                ],
                [
                    south_west,
                    {"lat": center_lat, "lng": min_lng - margin},
                    north_west,
                ],
                [
                    south_east,
                    {"lat": center_lat, "lng": max_lng + margin},
                    north_east,
                ],
            ]
        )

    def prepare(points: list[dict[str, float]]) -> list[dict[str, float]]:
        return sorted(
            points,
            key=lambda point: _route_projection_fraction(point, start, end),
        )

    options = [prepare(points) for points in options]
    safe_options = [
        points
        for points in options
        if not any(
            repository.contains_point(point) for point in points
        )
    ]
    options = safe_options or options
    clear_options = [
        points
        for points in options
        if not any(
            repository.segment_intersects_zone(a, b, zone)
            for a, b in zip([start, *points], [*points, end])
        )
    ]
    options = clear_options or options
    return sorted(options, key=lambda points: _path_distance([start, *points, end]))


def detour_waypoint_candidates(
    start: dict[str, float],
    end: dict[str, float],
    hazards: list[dict[str, Any]],
    limit: int = 8,
    mode: str = "pedestrian",
    repository: Any | None = None,
) -> list[list[dict[str, float]]]:
    repository = repository or get_flood_repository()
    ordered_hazards = sorted(
        hazards,
        key=lambda zone: _route_projection_fraction(_zone_center(zone), start, end),
    )
    option_sets = [
        _zone_detour_options(start, end, zone, mode, repository)
        for zone in ordered_hazards
    ]
    if not option_sets:
        return []

    corridor_candidates: list[
        tuple[tuple[int, float], list[dict[str, float]]]
    ] = []
    boundary_candidates: list[
        tuple[tuple[int, float], list[dict[str, float]]]
    ] = []
    seen: set[tuple[tuple[float, float], ...]] = set()

    def add_candidate(
        points: list[dict[str, float]],
        target: list[tuple[tuple[int, float], list[dict[str, float]]]],
    ) -> None:
        points = sorted(
            points,
            key=lambda point: _route_projection_fraction(point, start, end),
        )
        signature = tuple(
            (round(point["lat"], 7), round(point["lng"], 7))
            for point in points
        )
        if signature in seen or any(
            repository.contains_point(point) for point in points
        ):
            return
        seen.add(signature)
        path = [start, *points, end]
        intersections = sum(
            any(
                repository.segment_intersects_zone(a, b, zone)
                for a, b in zip(path, path[1:])
            )
            for zone in ordered_hazards
        )
        target.append(((intersections, _path_distance(path)), points))

    origin_lat = (start["lat"] + end["lat"]) / 2
    lng_scale = 111_320 * math.cos(math.radians(origin_lat))
    lat_scale = 110_574
    dx_m = (end["lng"] - start["lng"]) * lng_scale
    dy_m = (end["lat"] - start["lat"]) * lat_scale
    direct_distance = math.hypot(dx_m, dy_m)
    if direct_distance:
        perpendicular_x = -dy_m / direct_distance
        perpendicular_y = dx_m / direct_distance
        base_margin_m = max(
            float(zone.get("detour_margin_deg", 0.0015)) * lat_scale
            for zone in ordered_hazards
        )
        if mode == "pedestrian":
            offset_distances = [
                max(90, base_margin_m * 2 / 3),
                max(140, base_margin_m),
            ]
        else:
            offset_distances = [
                max(450, base_margin_m * 1.5),
                max(700, base_margin_m * 2),
            ]

        for fractions in ((0.05, 0.5, 0.95), (0.05, 0.35, 0.65, 0.95)):
            for offset_m in offset_distances:
                for side in (-1, 1):
                    points = [
                        {
                            "lat": start["lat"]
                            + (end["lat"] - start["lat"]) * fraction
                            + side * perpendicular_y * offset_m / lat_scale,
                            "lng": start["lng"]
                            + (end["lng"] - start["lng"]) * fraction
                            + side * perpendicular_x * offset_m / lng_scale,
                        }
                        for fraction in fractions
                    ]
                    add_candidate(points, corridor_candidates)

    for choices in product(*option_sets):
        point_variants: list[list[dict[str, float]]] = []
        flattened = [point for choice in choices for point in choice]
        if len(flattened) <= 5:
            point_variants.append(flattened)
        elif len(choices) == 2:
            point_variants.extend(
                [
                    [*choices[0], choices[1][0], choices[1][-1]],
                    [choices[0][0], choices[0][-1], *choices[1]],
                ]
            )
        else:
            point_variants.append([choice[len(choice) // 2] for choice in choices])

        for variant in point_variants:
            add_candidate(variant, boundary_candidates)

    corridor_candidates.sort(key=lambda item: item[0])
    boundary_candidates.sort(key=lambda item: item[0])
    candidates = []
    for index in range(max(len(corridor_candidates), len(boundary_candidates))):
        if index < len(boundary_candidates):
            candidates.append(boundary_candidates[index])
        if index < len(corridor_candidates):
            candidates.append(corridor_candidates[index])
    return [points for _, points in candidates[:limit]]


def detour_waypoints(
    start: dict[str, float],
    end: dict[str, float],
    hazards: list[dict[str, Any]],
    mode: str = "pedestrian",
    repository: Any | None = None,
) -> list[dict[str, float]]:
    candidates = detour_waypoint_candidates(
        start, end, hazards, limit=1, mode=mode, repository=repository
    )
    return candidates[0] if candidates else []


def annotate_shelter(
    shelter: dict[str, Any],
    origin: dict[str, float] | None,
    include_prediction: bool = False,
    repository: Any | None = None,
) -> dict[str, Any]:
    item = dict(shelter)
    if origin is None:
        item["distanceM"] = None
        item["riskZones"] = []
        item["score"] = None
        return item

    destination = {"lat": float(shelter["lat"]), "lng": float(shelter["lng"])}
    risks = (
        repository.hazards_for_paths([[origin, destination]])
        if repository is not None
        else hazards_for_segment(origin, destination, include_prediction)
    )
    distance = haversine_m(
        origin["lat"], origin["lng"], destination["lat"], destination["lng"]
    )
    item["distanceM"] = round(distance)
    item["riskZones"] = [{"id": risk["id"], "name": risk["name"]} for risk in risks]
    item["score"] = round(distance + 1_200 * len(risks))
    return item


def parse_float(value: str | None, field_name: str) -> float:
    try:
        return float(value or "")
    except ValueError as exc:
        raise ValueError(f"{field_name} 값이 올바르지 않습니다.") from exc


def _fallback_route(
    start: dict[str, float],
    end: dict[str, float],
    pass_points: list[dict[str, float]],
    mode: str,
) -> dict[str, Any]:
    path = [start, *pass_points, end]
    distance = sum(
        haversine_m(path[index]["lat"], path[index]["lng"], path[index + 1]["lat"], path[index + 1]["lng"])
        for index in range(len(path) - 1)
    )
    meters_per_second = 1.2 if mode == "pedestrian" else 9.5
    return {
        "source": "local-fallback",
        "summary": {
            "totalDistance": round(distance),
            "totalTime": round(distance / meters_per_second),
        },
        "path": path,
        "pathSegments": [path],
        "steps": [
            {
                "description": "로컬 대체 경로입니다. 티맵 API 키를 설정하면 실제 도로망 경로가 표시됩니다.",
                "distance": round(distance),
            }
        ],
    }


def _extract_tmap_route(payload: dict[str, Any]) -> dict[str, Any]:
    features = payload.get("features", [])
    path: list[dict[str, float]] = []
    path_segments: list[list[dict[str, float]]] = []
    steps = []
    summary: dict[str, Any] = {"totalDistance": None, "totalTime": None}

    for feature in features:
        geometry = feature.get("geometry", {})
        properties = feature.get("properties", {})
        coordinates = geometry.get("coordinates", [])

        if properties.get("totalDistance") is not None:
            summary["totalDistance"] = properties.get("totalDistance")
        if properties.get("totalTime") is not None:
            summary["totalTime"] = properties.get("totalTime")

        if geometry.get("type") == "Point" and properties.get("description"):
            steps.append(
                {
                    "description": properties.get("description"),
                    "distance": properties.get("distance"),
                    "turnType": properties.get("turnType"),
                }
            )

    line_features = [
        feature
        for feature in features
        if feature.get("geometry", {}).get("type") == "LineString"
    ]
    indexed_line_features = [
        feature
        for feature in line_features
        if feature.get("properties", {}).get("index") is not None
    ]
    route_line_features = indexed_line_features or line_features

    def feature_order(feature: dict[str, Any]) -> float:
        try:
            return float(feature.get("properties", {}).get("index"))
        except (TypeError, ValueError):
            return math.inf

    for feature in sorted(route_line_features, key=feature_order):
        coordinates = feature.get("geometry", {}).get("coordinates", [])
        segment = [
            {"lng": float(lng), "lat": float(lat)} for lng, lat in coordinates
        ]
        if not segment:
            continue
        if path:
            distance_to_start = haversine_m(
                path[-1]["lat"],
                path[-1]["lng"],
                segment[0]["lat"],
                segment[0]["lng"],
            )
            distance_to_end = haversine_m(
                path[-1]["lat"],
                path[-1]["lng"],
                segment[-1]["lat"],
                segment[-1]["lng"],
            )
            if distance_to_end < distance_to_start:
                segment.reverse()
        path_segments.append(segment)
        if path and path[-1] == segment[0]:
            path.extend(segment[1:])
        else:
            path.extend(segment)

    return {
        "source": "tmap",
        "summary": summary,
        "path": path,
        "pathSegments": path_segments,
        "steps": steps[:8],
        "rawFeatureCount": len(features),
    }


def request_tmap_route(
    start: dict[str, float],
    end: dict[str, float],
    end_name: str,
    pass_points: list[dict[str, float]],
    mode: str,
) -> tuple[dict[str, Any] | None, str | None]:
    app_key = os.getenv("TMAP_APP_KEY", "").strip()
    if not app_key:
        return None, "TMAP_APP_KEY가 설정되지 않았습니다."

    if mode == "car":
        endpoint = "https://apis.openapi.sk.com/tmap/routes"
        payload: dict[str, Any] = {
            "startX": start["lng"],
            "startY": start["lat"],
            "endX": end["lng"],
            "endY": end["lat"],
            "reqCoordType": "WGS84GEO",
            "resCoordType": "WGS84GEO",
            "startName": quote("현재 위치"),
            "endName": quote(end_name),
            "searchOption": 0,
            "carType": 0,
        }
    else:
        endpoint = "https://apis.openapi.sk.com/tmap/routes/pedestrian"
        payload = {
            "startX": start["lng"],
            "startY": start["lat"],
            "endX": end["lng"],
            "endY": end["lat"],
            "reqCoordType": "WGS84GEO",
            "resCoordType": "WGS84GEO",
            "startName": quote("현재 위치"),
            "endName": quote(end_name),
            "searchOption": 30,
            "speed": 4,
        }

    if pass_points:
        payload["passList"] = "_".join(
            f"{point['lng']},{point['lat']}" for point in pass_points[:5]
        )

    headers = {
        "Accept": "application/json",
        "Content-Type": "application/json",
        "appKey": app_key,
    }

    try:
        response = requests.post(
            endpoint,
            params={"version": "1", "format": "json"},
            headers=headers,
            json=payload,
            timeout=10,
        )
        response.raise_for_status()
        return _extract_tmap_route(response.json()), None
    except requests.RequestException as exc:
        return None, f"티맵 경로 API 호출 실패: {exc}"
    except (KeyError, TypeError, ValueError) as exc:
        return None, f"티맵 응답 처리 실패: {exc}"


@app.get("/")
def index() -> str:
    return render_template(
        "index.html",
        tmap_app_key=os.getenv("TMAP_APP_KEY", "").strip(),
    )


@app.get("/admin/reports")
def admin_reports() -> str:
    return render_template(
        "admin_reports.html",
        tmap_app_key=os.getenv("TMAP_APP_KEY", "").strip(),
    )


@app.get("/api/config")
def api_config():
    app_key = os.getenv("TMAP_APP_KEY", "").strip()
    return jsonify(
        {
            "tmapAppKey": app_key,
            "hasTmapKey": bool(app_key),
            "defaultCenter": {"lat": 36.1195, "lng": 128.3446},
        }
    )


@app.get("/api/flood-zones")
def api_flood_zones():
    return jsonify(load_flood_display())


@app.get("/api/live-flood-zones")
def api_live_flood_zones():
    return jsonify(live_flood_feature_collection(load_active_location_reports()))


@app.post("/api/reports")
def api_create_report():
    payload = request.get_json(silent=True) or {}
    try:
        lat = parse_float(str(payload.get("lat")), "lat")
        lng = parse_float(str(payload.get("lng")), "lng")
    except ValueError as exc:
        return jsonify({"error": str(exc)}), 400

    if not (
        REPORT_BOUNDS["minLat"] <= lat <= REPORT_BOUNDS["maxLat"]
        and REPORT_BOUNDS["minLng"] <= lng <= REPORT_BOUNDS["maxLng"]
    ):
        return jsonify({"error": "구미시 인근 좌표만 신고할 수 있습니다."}), 400

    accuracy_m = None
    if payload.get("accuracyM") is not None:
        try:
            accuracy_m = parse_float(str(payload.get("accuracyM")), "accuracyM")
        except ValueError as exc:
            return jsonify({"error": str(exc)}), 400
        if not 0 <= accuracy_m <= 10_000:
            return jsonify({"error": "accuracyM 값이 허용 범위를 벗어났습니다."}), 400

    note = str(payload.get("note") or "").strip()
    if len(note) > REPORT_NOTE_MAX_LENGTH:
        return jsonify(
            {"error": f"신고 메모는 {REPORT_NOTE_MAX_LENGTH}자 이하로 입력해 주세요."}
        ), 400

    report = save_location_report(lat, lng, accuracy_m, note)
    return jsonify({"report": report}), 201


@app.get("/api/admin/reports")
def api_admin_reports():
    reports = load_location_reports()
    confirmed_count = sum(
        report["status"] in {"active", "resolved"} for report in reports
    )
    return jsonify(
        {
            "reports": reports,
            "count": len(reports),
            "confirmedCount": confirmed_count,
        }
    )


@app.post("/api/admin/flood-zones")
def api_create_admin_flood_zone():
    payload = request.get_json(silent=True) or {}
    try:
        lat = parse_float(str(payload.get("lat")), "lat")
        lng = parse_float(str(payload.get("lng")), "lng")
    except ValueError as exc:
        return jsonify({"error": str(exc)}), 400

    if not (
        REPORT_BOUNDS["minLat"] <= lat <= REPORT_BOUNDS["maxLat"]
        and REPORT_BOUNDS["minLng"] <= lng <= REPORT_BOUNDS["maxLng"]
    ):
        return jsonify({"error": "구미시 인근 좌표만 추가할 수 있습니다."}), 400

    note = str(payload.get("note") or "").strip()
    if len(note) > REPORT_NOTE_MAX_LENGTH:
        return jsonify(
            {"error": f"현장 메모는 {REPORT_NOTE_MAX_LENGTH}자 이하로 입력해 주세요."}
        ), 400

    report = save_location_report(
        lat,
        lng,
        None,
        note or "관리자 지도 직접 추가",
        source="admin",
    )
    report = update_location_report_status(report["id"], "approve")
    return jsonify({"report": report}), 201


@app.get("/api/admin/reports/export.xlsx")
def api_export_reports():
    reports = load_confirmed_flood_reports()
    output = build_flood_history_workbook(reports)
    filename = f"gumi_flood_history_{datetime.now(KST):%Y%m%d}.xlsx"
    return send_file(
        output,
        as_attachment=True,
        download_name=filename,
        mimetype=(
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        ),
    )


@app.post("/api/admin/reports/<int:report_id>/approve")
def api_approve_report(report_id: int):
    report = update_location_report_status(report_id, "approve")
    if report is None:
        return jsonify({"error": "신고를 찾을 수 없습니다."}), 404
    return jsonify({"report": report})


@app.post("/api/admin/reports/<int:report_id>/resolve")
def api_resolve_report(report_id: int):
    try:
        report = update_location_report_status(report_id, "resolve")
    except ValueError as exc:
        return jsonify({"error": str(exc)}), 409
    if report is None:
        return jsonify({"error": "신고를 찾을 수 없습니다."}), 404
    return jsonify({"report": report})


@app.get("/api/shelters")
def api_shelters():
    lat = request.args.get("lat")
    lng = request.args.get("lng")
    origin = None
    if lat is not None and lng is not None:
        try:
            origin = {"lat": parse_float(lat, "lat"), "lng": parse_float(lng, "lng")}
        except ValueError as exc:
            return jsonify({"error": str(exc)}), 400

    include_history = request.args.get("includeRiskAreas", "").lower() in {
        "1",
        "true",
        "yes",
    }
    include_prediction = request.args.get("includeFiftyYear", "").lower() in {
        "1",
        "true",
        "yes",
    }
    include_live_flood = True
    flood_repository = flood_repository_for_request(
        include_prediction, include_live_flood, include_history
    )
    shelters = [
        annotate_shelter(
            shelter,
            origin,
            include_prediction,
            repository=flood_repository,
        )
        for shelter in load_shelters()
    ]
    if origin is not None:
        shelters.sort(key=lambda shelter: (shelter["score"], shelter["distanceM"]))

    if shelters:
        data_notice = (
            "재난안전데이터 공유 플랫폼의 공식 구미시 이재민 임시주거시설 "
            f"{len(shelters)}곳을 제공합니다."
        )
    else:
        data_notice = (
            "공식 이재민 임시주거시설 데이터가 비어 있습니다. "
            "DSSP-IF-10945 이용신청과 서비스 키 설정을 확인해 주세요."
        )

    return jsonify(
        {
            "shelters": shelters,
            "nearest": shelters[0] if shelters else None,
            "dataNotice": data_notice,
        }
    )


@app.post("/api/route")
def api_route():
    payload = request.get_json(silent=True) or {}
    start = payload.get("start") or {}
    shelter_id = payload.get("shelterId")
    mode = payload.get("mode", "pedestrian")
    avoid_flood = bool(payload.get("avoidFlood", True))
    include_history = bool(payload.get("includeRiskAreas", False))
    include_prediction = bool(payload.get("includeFiftyYear", False))
    include_live_flood = True

    if mode not in {"pedestrian", "car"}:
        return jsonify({"error": "mode는 pedestrian 또는 car만 가능합니다."}), 400

    try:
        start_point = {
            "lat": parse_float(str(start.get("lat")), "start.lat"),
            "lng": parse_float(str(start.get("lng")), "start.lng"),
        }
    except ValueError as exc:
        return jsonify({"error": str(exc)}), 400

    shelter = next(
        (item for item in load_shelters() if item["id"] == shelter_id),
        None,
    )
    if shelter is None:
        return jsonify({"error": "대피소를 찾을 수 없습니다."}), 404

    end_point = {"lat": float(shelter["lat"]), "lng": float(shelter["lng"])}
    route, tmap_error = request_tmap_route(
        start_point, end_point, shelter["name"], [], mode
    )
    route_attempts = 1
    pass_points: list[dict[str, float]] = []
    warning_parts: list[str] = []
    initial_risk_distance_m = 0.0
    flood_repository = flood_repository_for_request(
        include_prediction, include_live_flood, include_history
    )
    if (
        avoid_flood
        and include_history
        and getattr(flood_repository, "is_empty", False)
    ):
        warning_parts.append(
            "구미시 공식 침수흔적 등록 건이 없어 우회 판정에 반영할 구역이 없습니다. 기록 없음이 침수 위험 없음을 의미하지는 않습니다."
        )
    unavoidable_ids = {
        zone["id"]
        for zone in [
            *flood_repository.hazards_for_point(start_point),
            *flood_repository.hazards_for_point(end_point),
        ]
    }
    access_limited_ids: set[str] = set()
    nearby_hazards = getattr(flood_repository, "hazards_near_point", None)
    if mode == "car" and callable(nearby_hazards):
        access_limited_ids = {
            zone["id"] for zone in nearby_hazards(end_point, 150)
        } - unavoidable_ids

    if route is None:
        initial_hazards = flood_repository.hazards_for_paths(
            [[start_point, end_point]]
        )
        avoidable_hazards = [
            zone for zone in initial_hazards if zone["id"] not in unavoidable_ids
        ]
        if avoid_flood and avoidable_hazards:
            pass_points = detour_waypoints(
                start_point,
                end_point,
                avoidable_hazards,
                mode=mode,
                repository=flood_repository,
            )
        route = _fallback_route(start_point, end_point, pass_points, mode)
        initial_risk_distance_m = _route_risk_distance_m(route, flood_repository)
        if tmap_error:
            warning_parts.append(tmap_error)
    else:
        initial_hazards = hazards_for_route(route, flood_repository)
        initial_risk_distance_m = _route_risk_distance_m(route, flood_repository)
        avoidable_hazards = [
            zone for zone in initial_hazards if zone["id"] not in unavoidable_ids
        ]

        if avoid_flood and avoidable_hazards:
            best_route = route
            best_pass_points: list[dict[str, float]] = []
            best_hazards = initial_hazards

            def route_score(
                candidate_route: dict[str, Any], candidate_hazards: list[dict[str, Any]]
            ) -> tuple[float, int, float]:
                remaining_count = sum(
                    zone["id"] not in unavoidable_ids for zone in candidate_hazards
                )
                risk_distance = _route_risk_distance_m(
                    candidate_route, flood_repository
                )
                distance = candidate_route.get("summary", {}).get("totalDistance")
                return risk_distance, remaining_count, float(distance or math.inf)

            best_score = route_score(best_route, best_hazards)
            candidates = detour_waypoint_candidates(
                start_point,
                end_point,
                avoidable_hazards,
                limit=8 if mode == "car" else 6,
                mode=mode,
                repository=flood_repository,
            )
            for candidate_points in candidates:
                candidate_route, _ = request_tmap_route(
                    start_point,
                    end_point,
                    shelter["name"],
                    candidate_points,
                    mode,
                )
                route_attempts += 1
                if candidate_route is None:
                    continue

                candidate_hazards = hazards_for_route(
                    candidate_route, flood_repository
                )
                candidate_score = route_score(candidate_route, candidate_hazards)
                if candidate_score < best_score:
                    best_route = candidate_route
                    best_pass_points = candidate_points
                    best_hazards = candidate_hazards
                    best_score = candidate_score

            route = best_route
            pass_points = best_pass_points

    final_hazards = hazards_for_route(route, flood_repository)
    final_risk_distance_m = _route_risk_distance_m(route, flood_repository)
    risk_reduction_percent = 0
    if initial_risk_distance_m > 0 and final_risk_distance_m < initial_risk_distance_m:
        risk_reduction_percent = round(
            (1 - final_risk_distance_m / initial_risk_distance_m) * 100
        )
    remaining_avoidable = [
        zone
        for zone in final_hazards
        if zone["id"] not in unavoidable_ids | access_limited_ids
    ]
    remaining_access_limited = [
        zone for zone in final_hazards if zone["id"] in access_limited_ids
    ]
    remaining_unavoidable = [
        zone for zone in final_hazards if zone["id"] in unavoidable_ids
    ]
    final_ids = {zone["id"] for zone in final_hazards}
    avoided_hazards = [
        zone for zone in initial_hazards if zone["id"] not in final_ids
    ]

    if avoid_flood and remaining_avoidable:
        if risk_reduction_percent >= 5:
            warning_parts.append(
                "완전 우회 가능한 경로를 찾지 못해 침수 영향 구간을 "
                f"약 {round(initial_risk_distance_m)}m에서 "
                f"{round(final_risk_distance_m)}m로 최소화했습니다."
            )
        else:
            warning_parts.append(
                f"도로망 제약으로 침수 위험지역 {len(remaining_avoidable)}곳을 완전히 우회하지 못했습니다."
            )
    if avoid_flood and remaining_access_limited:
        warning_parts.append(
            "대피소 인근 진입도로가 침수 예상구역에 포함되어 해당 구간은 완전히 피할 수 없습니다."
        )
    if avoid_flood and remaining_unavoidable:
        warning_parts.append(
            "출발지 또는 대피소가 침수 위험지역 안에 있어 해당 구역은 완전히 피할 수 없습니다."
        )

    if not avoid_flood:
        avoidance_status = "disabled"
    elif remaining_avoidable:
        avoidance_status = "partial"
    elif remaining_unavoidable or remaining_access_limited:
        avoidance_status = "unavoidable"
    else:
        avoidance_status = "clear"

    route["shelter"] = shelter
    route["detectedRiskZones"] = [
        {"id": zone["id"], "name": zone["name"]} for zone in initial_hazards
    ]
    route["riskZones"] = [
        {"id": zone["id"], "name": zone["name"]} for zone in final_hazards
    ]
    route["avoidedRiskZones"] = [
        {"id": zone["id"], "name": zone["name"]} for zone in avoided_hazards
    ]
    route["accessLimitedRiskZones"] = [
        {"id": zone["id"], "name": zone["name"]}
        for zone in remaining_access_limited
    ]
    route["passPoints"] = pass_points
    route["avoidFlood"] = avoid_flood
    route["includeRiskAreas"] = include_history
    route["includeFiftyYear"] = include_prediction
    route["includeLiveFlood"] = include_live_flood
    route["avoidanceStatus"] = avoidance_status
    route["routeAttempts"] = route_attempts
    route["initialRiskDistanceM"] = round(initial_risk_distance_m)
    route["riskDistanceM"] = round(final_risk_distance_m)
    route["riskReductionPercent"] = risk_reduction_percent
    if warning_parts:
        route["warning"] = " ".join(warning_parts)
    return jsonify(route)


if __name__ == "__main__":
    app.run(host="0.0.0.0", port=int(os.getenv("PORT", "5000")), debug=True)
